from django.db.models import Count, Sum, Q, Avg
from django.utils import timezone
from datetime import datetime, timedelta
from decimal import Decimal

from account.models import User
from estate.models import Customer, Property, Match
from case_management.models import CustomerProcess, Visit, SupervisorAlert
from transaction.models import Contract, Commission, WithdrawalRequest


def get_dashboard_stats():
    """دریافت آمار کلی برای داشبورد"""
    now = timezone.now()
    today_start = now.replace(hour=0, minute=0, second=0, microsecond=0)

    stats = {
        # مشتریان
        'total_customers': Customer.objects.filter(is_active=True).count(),
        'new_customers_today': Customer.objects.filter(created_at__gte=today_start).count(),
        'customers_in_meeting': Customer.objects.filter(status=Customer.Status.IN_MEETING).count(),

        # املاک
        'total_properties': Property.objects.filter(is_active=True).count(),
        'new_properties_today': Property.objects.filter(created_at__gte=today_start).count(),
        'properties_in_meeting': Property.objects.filter(status=Property.Status.IN_MEETING).count(),

        # کاربران
        'total_experts': User.objects.filter(user_roles__role__name='expert', is_active=True).distinct().count(),
        'total_supervisors': User.objects.filter(user_roles__role__name='supervisor', is_active=True).distinct().count(),

        # فرآیندها
        'total_processes': CustomerProcess.objects.count(),
        'active_processes': CustomerProcess.objects.exclude(
            status__in=[CustomerProcess.Status.CLOSED_SUCCESS,
                       CustomerProcess.Status.CLOSED_FAIL,
                       CustomerProcess.Status.CANCELLED,
                       CustomerProcess.Status.CONVERTED,
                       CustomerProcess.Status.SENT_TO_ACCOUNTING]
        ).count(),
        'closed_success': CustomerProcess.objects.filter(status=CustomerProcess.Status.CLOSED_SUCCESS).count(),
        'closed_fail': CustomerProcess.objects.filter(status=CustomerProcess.Status.CLOSED_FAIL).count(),
        'contract_pending': CustomerProcess.objects.filter(status=CustomerProcess.Status.CONTRACT_PENDING).count(),
        'in_accounting': CustomerProcess.objects.filter(status=CustomerProcess.Status.SENT_TO_ACCOUNTING).count(),
        'alerts_unread': SupervisorAlert.objects.filter(is_read=False).count(),
        'missed_deadlines': SupervisorAlert.objects.filter(
            is_read=False,
            alert_type__in=[SupervisorAlert.Type.MISSED_CONTACT,
                           SupervisorAlert.Type.MISSED_INTRODUCTION,
                           SupervisorAlert.Type.MISSED_VISIT_RESULT]
        ).count(),

        # قراردادها
        'total_contracts': Contract.objects.count(),
        'completed_contracts': Contract.objects.filter(status=Contract.Status.COMPLETED).count(),
        'pending_contracts': Contract.objects.filter(status=Contract.Status.PENDING).count(),

        # حسابداری
        'total_commissions': Commission.objects.count(),
        'paid_commissions': Commission.objects.filter(status=Commission.Status.PAID).count(),
        'pending_commissions': Commission.objects.filter(status=Commission.Status.PENDING).count(),
        'wallet_transferred': Commission.objects.filter(status=Commission.Status.WALLET_TRANSFERRED).count(),
        'total_commission_amount': Commission.objects.aggregate(Sum('total_commission'))['total_commission__sum'] or 0,

        # درخواست‌های برداشت
        'withdrawal_pending': WithdrawalRequest.objects.filter(status=WithdrawalRequest.Status.PENDING).count(),
        'withdrawal_approved': WithdrawalRequest.objects.filter(status=WithdrawalRequest.Status.APPROVED).count(),
        'withdrawal_paid': WithdrawalRequest.objects.filter(status=WithdrawalRequest.Status.PAID).count(),
    }

    return stats


def get_status_chart_data():
    """داده‌های نمودار وضعیت‌ها"""
    # وضعیت مشتریان
    customer_status = Customer.objects.filter(is_active=True).values('status').annotate(count=Count('id'))
    customer_status_dict = {item['status']: item['count'] for item in customer_status}

    # وضعیت املاک
    property_status = Property.objects.filter(is_active=True).values('status').annotate(count=Count('id'))
    property_status_dict = {item['status']: item['count'] for item in property_status}

    # وضعیت فرآیندها
    process_status = CustomerProcess.objects.values('status').annotate(count=Count('id'))
    process_status_dict = {item['status']: item['count'] for item in process_status}

    return {
        'customer': customer_status_dict,
        'property': property_status_dict,
        'process': process_status_dict,
    }


def get_expert_performance():
    """گزارش عملکرد کارشناسان"""
    experts = User.objects.filter(
        user_roles__role__name='expert',
        is_active=True
    ).distinct()

    result = []
    for expert in experts:
        processes = CustomerProcess.objects.filter(expert=expert)
        total = processes.count()
        closed = processes.filter(
            status__in=[CustomerProcess.Status.CLOSED_SUCCESS,
                       CustomerProcess.Status.CLOSED_FAIL,
                       CustomerProcess.Status.CANCELLED,
                       CustomerProcess.Status.CONVERTED,
                       CustomerProcess.Status.SENT_TO_ACCOUNTING]
        ).count()
        success = processes.filter(status=CustomerProcess.Status.CLOSED_SUCCESS).count()
        fail = processes.filter(status=CustomerProcess.Status.CLOSED_FAIL).count()
        converted = processes.filter(status=CustomerProcess.Status.CONVERTED).count()
        active = total - closed

        # محاسبه نرخ موفقیت
        if closed > 0:
            success_rate = round((success / closed) * 100, 2)
        else:
            success_rate = 0

        # میانگین زمان بسته شدن
        closed_processes = processes.filter(closed_at__isnull=False)
        avg_days = 0
        if closed_processes.exists():
            total_days = 0
            for p in closed_processes:
                days = (p.closed_at - p.assigned_at).days
                total_days += days
            avg_days = round(total_days / closed_processes.count(), 2)

        result.append({
            'expert': expert,
            'total': total,
            'active': active,
            'closed': closed,
            'success': success,
            'fail': fail,
            'converted': converted,
            'success_rate': success_rate,
            'avg_close_days': avg_days,
        })

    return sorted(result, key=lambda x: x['success_rate'], reverse=True)


def get_supervisor_stats():
    """آمار سرپرستان و کارشناسان زیرمجموعه"""
    supervisors = User.objects.filter(
        user_roles__role__name='supervisor',
        is_active=True
    ).distinct()

    result = []
    for supervisor in supervisors:
        experts = User.objects.filter(referral=supervisor, is_active=True)
        experts_count = experts.count()

        # فرآیندهای زیرمجموعه
        processes = CustomerProcess.objects.filter(supervisor=supervisor)
        total = processes.count()
        active = processes.exclude(
            status__in=[CustomerProcess.Status.CLOSED_SUCCESS,
                       CustomerProcess.Status.CLOSED_FAIL,
                       CustomerProcess.Status.CANCELLED,
                       CustomerProcess.Status.CONVERTED,
                       CustomerProcess.Status.SENT_TO_ACCOUNTING]
        ).count()
        success = processes.filter(status=CustomerProcess.Status.CLOSED_SUCCESS).count()

        result.append({
            'supervisor': supervisor,
            'experts_count': experts_count,
            'experts': experts,
            'total_processes': total,
            'active_processes': active,
            'success_processes': success,
        })

    return result


def get_monthly_stats():
    """آمار ماهانه برای نمودارها"""
    now = timezone.now()
    months = []
    labels = []

    # ۱۲ ماه گذشته
    for i in range(11, -1, -1):
        month_start = (now - timedelta(days=30 * i)).replace(day=1, hour=0, minute=0, second=0, microsecond=0)
        if i == 0:
            month_end = now
        else:
            next_month = (month_start + timedelta(days=32)).replace(day=1)
            month_end = next_month

        labels.append(month_start.strftime('%Y/%m'))

        # تعداد مشتریان جدید
        customers = Customer.objects.filter(created_at__gte=month_start, created_at__lt=month_end).count()

        # تعداد املاک جدید
        properties = Property.objects.filter(created_at__gte=month_start, created_at__lt=month_end).count()

        # تعداد قراردادهای جدید
        contracts = Contract.objects.filter(created_at__gte=month_start, created_at__lt=month_end).count()

        months.append({
            'label': month_start.strftime('%Y/%m'),
            'customers': customers,
            'properties': properties,
            'contracts': contracts,
        })

    return months


def get_contract_type_stats():
    """آمار قراردادها بر اساس نوع"""
    contract_types = Contract.CONTRACT_TYPE_CHOICES
    result = []
    total = Contract.objects.count()

    for code, label in contract_types:
        count = Contract.objects.filter(contract_type=code).count()
        amount = Contract.objects.filter(contract_type=code).aggregate(Sum('final_amount'))['final_amount__sum'] or 0
        percentage = round((count / total * 100), 2) if total > 0 else 0
        result.append({
            'code': code,
            'label': label,
            'count': count,
            'amount': amount,
            'percentage': percentage,
        })

    return result


def get_commission_stats():
    """آمار کمیسیون‌ها"""
    total_commission = Commission.objects.aggregate(Sum('total_commission'))['total_commission__sum'] or 0
    total_paid = Commission.objects.filter(status=Commission.Status.PAID).aggregate(Sum('total_commission'))['total_commission__sum'] or 0
    total_pending = Commission.objects.filter(status=Commission.Status.PENDING).aggregate(Sum('total_commission'))['total_commission__sum'] or 0
    total_calculated = Commission.objects.filter(status=Commission.Status.CALCULATED).aggregate(Sum('total_commission'))['total_commission__sum'] or 0
    total_wallet = Commission.objects.filter(status=Commission.Status.WALLET_TRANSFERRED).aggregate(Sum('total_commission'))['total_commission__sum'] or 0

    return {
        'total': total_commission,
        'paid': total_paid,
        'pending': total_pending,
        'calculated': total_calculated,
        'wallet_transferred': total_wallet,
        'total_count': Commission.objects.count(),
        'paid_count': Commission.objects.filter(status=Commission.Status.PAID).count(),
        'pending_count': Commission.objects.filter(status=Commission.Status.PENDING).count(),
        'calculated_count': Commission.objects.filter(status=Commission.Status.CALCULATED).count(),
        'wallet_count': Commission.objects.filter(status=Commission.Status.WALLET_TRANSFERRED).count(),
    }


def export_to_excel(data, filename, columns, sheet_name='Sheet1'):
    """
    خروجی Excel از داده‌ها
    نیاز به نصب: pip install openpyxl
    """
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, PatternFill
    from openpyxl.utils import get_column_letter
    from django.http import HttpResponse

    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name

    # هدرها
    header_fill = PatternFill(start_color='366092', end_color='366092', fill_type='solid')
    header_font = Font(color='FFFFFF', bold=True)

    for col, (key, label) in enumerate(columns.items(), 1):
        cell = ws.cell(row=1, column=col, value=label)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal='center')

    # داده‌ها
    for row_idx, item in enumerate(data, 2):
        for col_idx, key in enumerate(columns.keys(), 1):
            value = item.get(key)
            if value is None:
                value = ''
            ws.cell(row=row_idx, column=col_idx, value=value)

    # تنظیم عرض ستون‌ها
    for col in range(1, len(columns) + 1):
        ws.column_dimensions[get_column_letter(col)].width = 20

    # تنظیم پاسخ
    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = f'attachment; filename={filename}.xlsx'
    wb.save(response)
    return response