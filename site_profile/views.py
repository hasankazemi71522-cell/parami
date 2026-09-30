import openpyxl
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
from django.http import HttpResponse
import os
from random import randint

import xlwt
from django.contrib.auth import get_user_model, logout, login
from django.contrib.auth.decorators import login_required, permission_required
from django.contrib.admin.views.decorators import staff_member_required
from django.core.exceptions import PermissionDenied
from django.db.models import Q
from django.http import JsonResponse
from django.shortcuts import render, redirect, get_object_or_404
from django.contrib import messages
from django.db import IntegrityError
from django.contrib.auth.hashers import make_password

from account.models import User, Role, UserRole, Permission, RolePermission
from account.myclass import send_message_login, new_income, send_message_signup
from estate.models import Customer, Property, Province, City
from parami import settings
from .other_class.profile_class import three_digits
from persiantools.jdatetime import JalaliDate
from myclass.mydef import logout_user, edit_back_mode, views_permissions
from myclass.mydef import to_shamsi_date, to_shamsi_datetime
from datetime import datetime
from django.utils import timezone

user = get_user_model()


# ============================================================
# 🗓️ توابع کمکی برای تبدیل تاریخ به رشته شمسی
# ============================================================

def get_shamsi_date(value):
    """تبدیل تاریخ میلادی به رشته شمسی (فقط تاریخ) - مثال: 1403/07/05"""
    if not value:
        return ''
    try:
        jalali = to_shamsi_date(value)
        if jalali:
            return f"{jalali.year}/{jalali.month:02d}/{jalali.day:02d}"
    except Exception as e:
        print(f"خطا در تبدیل تاریخ شمسی: {e}")
    return ''


def get_shamsi_datetime(value):
    """تبدیل تاریخ میلادی به رشته شمسی (تاریخ و ساعت) - مثال: 1403/07/05 14:30"""
    if not value:
        return ''
    try:
        jalali = to_shamsi_datetime(value)
        if jalali:
            return f"{jalali.year}/{jalali.month:02d}/{jalali.day:02d} {jalali.hour:02d}:{jalali.minute:02d}"
    except Exception as e:
        print(f"خطا در تبدیل تاریخ و ساعت شمسی: {e}")
    return ''


def get_shamsi_time(value):
    """تبدیل تاریخ میلادی به رشته شمسی (فقط ساعت) - مثال: 14:30"""
    if not value:
        return ''
    try:
        jalali = to_shamsi_datetime(value)
        if jalali:
            return f"{jalali.hour:02d}:{jalali.minute:02d}"
    except Exception as e:
        print(f"خطا در تبدیل ساعت شمسی: {e}")
    return ''


def is_ajax(request):
    return request.META.get('HTTP_X_REQUESTED_WITH') == 'XMLHttpRequest'


# ==================== دکوریتور ساده برای بررسی دسترسی ====================
def role_required(role_names):
    """بررسی می‌کند کاربر نقش مورد نظر را دارد یا ادمین است"""

    def decorator(view_func):
        def wrapper(request, *args, **kwargs):
            if not request.user.is_authenticated:
                return redirect('/login/')

            # ادمین به همه چیز دسترسی دارد
            if request.user.is_admin:
                return view_func(request, *args, **kwargs)

            # بررسی نقش کاربر
            if not request.user.has_any_role(role_names):
                return redirect('/page_404/')

            return view_func(request, *args, **kwargs)

        return wrapper

    return decorator


# ==================== دکوریتور دسترسی به پروفایل ====================
def profile_access():
    """کاربر فقط به پروفایل خودش دسترسی دارد، ادمین به همه"""

    def decorator(view_func):
        def wrapper(request, userId, *args, **kwargs):
            if not request.user.is_authenticated:
                return redirect('/login/')

            # ادمین به همه دسترسی دارد
            if request.user.is_admin:
                return view_func(request, userId, *args, **kwargs)

            # کاربر عادی فقط به خودش
            if request.user.id == int(userId):
                return view_func(request, userId, *args, **kwargs)

            return redirect('/page_404/')

        return wrapper

    return decorator


@login_required
def dashboard(request):
    user = request.user
    permission, perimissin_list, perimissin_group = views_permissions(user, "user_dashboard")

    # مشتریانی که user آنها برابر کاربر فعلی است
    customers = Customer.objects.filter(user=user, is_active=True).order_by('-created_at')

    # فایل‌هایی که owner آنها برابر کاربر فعلی است
    properties = Property.objects.filter(owner=user, is_active=True).order_by('-created_at')

    # ✅ اضافه کردن تاریخ شمسی به مشتریان
    for c in customers:
        c.created_at_shamsi = get_shamsi_date(c.created_at)
        c.updated_at_shamsi = get_shamsi_date(c.updated_at)

    # ✅ اضافه کردن تاریخ شمسی به املاک
    for p in properties:
        p.created_at_shamsi = get_shamsi_date(p.created_at)
        p.updated_at_shamsi = get_shamsi_date(p.updated_at)

    context = {
        'title': 'داشبورد کاربری',
        'userId': user.id,
        'this_user': user,
        'customers': customers,
        'properties': properties,
        "perimissin_list": perimissin_list,
        "perimissin_group": perimissin_group,
    }
    return render(request, 'dashboard.html', context)


@login_required
def page_404(request):
    user = request.user
    permission, perimissin_list, perimissin_group = views_permissions(user, "permission_delete")

    context = {
        "title": "خطای 404",
        "userId": request.user.id if request.user.is_authenticated else None,
        "this_user": request.user if request.user.is_authenticated else None,
        "perimissin_list": perimissin_list,
        "perimissin_group": perimissin_group,
    }
    return render(request, "404.html", context)


# ==================== تنظیمات پروفایل ====================

@login_required
def profile_settings(request, userId):
    this_user = User.objects.get(id=userId)
    user = request.user
    permission, perimissin_list, perimissin_group = views_permissions(user, "profile_settings")

    # دریافت لیست استان‌ها و شهرها
    provinces = Province.objects.all().order_by('name')
    cities = City.objects.select_related('province').all().order_by('province__name', 'name')

    if request.method == 'POST':
        first_name = request.POST.get('first_name')
        last_name = request.POST.get('last_name')
        email = request.POST.get('email')
        mobile = request.POST.get('mobile')
        phone = request.POST.get('phone')
        state = request.POST.get('state')
        city = request.POST.get('city')
        home_address = request.POST.get('home_address')
        post_code = request.POST.get('post_code')

        bank_name = request.POST.get('bank_name')
        bank_card_number = request.POST.get('bank_card_number')
        bank_account_number = request.POST.get('bank_account_number')
        bank_sheba_number = request.POST.get('bank_sheba_number')

        this_user.first_name = first_name
        this_user.last_name = last_name
        this_user.email = email
        this_user.mobile = mobile
        this_user.phone = phone
        this_user.state = state
        this_user.city = city
        this_user.home_address = home_address
        this_user.post_code = post_code

        this_user.bank_name = bank_name
        this_user.bank_card_number = bank_card_number
        this_user.bank_account_number = bank_account_number
        this_user.bank_sheba_number = bank_sheba_number

        if request.FILES.get('user_image'):
            this_user.user_image = request.FILES.get('user_image')

        this_user.save(update_fields=[
            'first_name', 'last_name', 'email', 'mobile', 'phone',
            'state', 'city', 'home_address', 'post_code',
            'bank_name', 'bank_card_number', 'bank_account_number', 'bank_sheba_number',
            'user_image'
        ])

        messages.success(request, 'اطلاعات شما با موفقیت بروزرسانی شد')
        return redirect('site_profile:profile_settings', userId=userId)

    # ✅ اضافه کردن تاریخ شمسی به این کاربر
    this_user.date_joined_shamsi = get_shamsi_datetime(this_user.date_joined)
    this_user.last_login_shamsi = get_shamsi_datetime(this_user.last_login)

    context = {
        'userId': userId,
        'title': 'تنظیمات پروفایل',
        'this_user': request.user,
        "perimissin_list": perimissin_list,
        "perimissin_group": perimissin_group,
        'provinces': provinces,
        'cities': cities,
    }
    return render(request, 'profile_settings.html', context)


# ==================== مدیریت کاربران ====================
@login_required
def user_list(request):
    """لیست همه کاربران با قابلیت جستجو و فیلتر"""
    user = request.user
    permission, perimissin_list, perimissin_group = views_permissions(user, "user_list")

    if permission is False:
        messages.error(request, 'شما اجازه استفاده از این فرم را ندارید')
        return redirect('site_profile:page_404')

    users = User.objects.all().order_by('-date_joined')

    # جستجو
    search = request.GET.get('search', '')
    if search:
        users = users.filter(
            Q(username__icontains=search) |
            Q(first_name__icontains=search) |
            Q(last_name__icontains=search) |
            Q(mobile__icontains=search) |
            Q(email__icontains=search)
        )

    # فیلتر بر اساس نقش
    role_filter = request.GET.get('role', '')
    if role_filter:
        users = users.filter(user_roles__role__name=role_filter, user_roles__is_active=True)

    # لیست نقش‌ها برای فیلتر
    roles = Role.objects.filter(is_active=True)

    # ✅ اضافه کردن تاریخ شمسی به هر کاربر
    for u in users:
        u.date_joined_shamsi = get_shamsi_date(u.date_joined)
        u.last_login_shamsi = get_shamsi_datetime(u.last_login)

    context = {
        'users': users,
        'roles': roles,
        'search': search,
        'role_filter': role_filter,
        'title': 'لیست کاربران',
        'userId': request.user.id,
        'this_user': request.user,
        "perimissin_list": perimissin_list,
        "perimissin_group": perimissin_group,
    }
    return render(request, 'users/user_list.html', context)


@login_required
def user_create(request):
    """افزودن کاربر جدید"""
    user = request.user
    permission, perimissin_list, perimissin_group = views_permissions(user, "user_create")

    if permission is False:
        messages.error(request, 'شما اجازه استفاده از این فرم را ندارید')
        return redirect('site_profile:page_404')

    roles = Role.objects.filter(is_active=True)

    if request.method == 'POST':
        username = request.POST.get('username')
        first_name = request.POST.get('first_name')
        last_name = request.POST.get('last_name')
        mobile = request.POST.get('mobile')
        email = request.POST.get('email')
        password = request.POST.get('password')
        confirm_password = request.POST.get('confirm_password')
        is_active = request.POST.get('is_active') == 'on'
        selected_roles = request.POST.getlist('roles')

        errors = []

        if not username:
            errors.append('نام کاربری الزامی است')
        elif User.objects.filter(username=username).exists():
            errors.append('این نام کاربری قبلاً ثبت شده است')

        if mobile and User.objects.filter(mobile=mobile).exists():
            errors.append('این شماره موبایل قبلاً ثبت شده است')

        if email and User.objects.filter(email=email).exists():
            errors.append('این ایمیل قبلاً ثبت شده است')

        if not password:
            errors.append('رمز عبور الزامی است')
        elif len(password) < 4:
            errors.append('رمز عبور باید حداقل ۴ کاراکتر باشد')
        elif password != confirm_password:
            errors.append('رمز عبور و تکرار آن مطابقت ندارند')

        if errors:
            for error in errors:
                messages.error(request, error)
            return render(request, 'users/user_create.html', {
                'roles': roles,
                'form_data': request.POST,
                'userId': request.user.id,
                'this_user': request.user,
            })

        try:
            user = User.objects.create(
                username=username,
                first_name=first_name,
                last_name=last_name,
                mobile=mobile,
                email=email,
                is_active=is_active,
                password=make_password(password)
            )

            # اختصاص نقش‌ها
            for role_name in selected_roles:
                try:
                    role = Role.objects.get(name=role_name)
                    UserRole.objects.create(
                        user=user,
                        role=role,
                        assigned_by=request.user,
                        is_active=True
                    )
                except Role.DoesNotExist:
                    pass

            full_name = user.first_name + " " + user.last_name

            send_message_signup(full_name, user.mobile, user.username, password)
            messages.success(request, f'کاربر "{user.get_full_name()}" با موفقیت ایجاد شد')
            return redirect('site_profile:user_list')

        except Exception as e:
            messages.error(request, f'خطا در ایجاد کاربر: {str(e)}')

    context = {
        'roles': roles,
        'title': 'افزودن کاربر جدید',
        'userId': request.user.id,
        'this_user': request.user,
        "perimissin_list": perimissin_list,
        "perimissin_group": perimissin_group,
    }
    return render(request, 'users/user_create.html', context)


@login_required
def user_edit(request, user_id):
    """ویرایش کاربر"""
    user = request.user
    permission, perimissin_list, perimissin_group = views_permissions(user, "user_edit")

    if permission is False:
        messages.error(request, 'شما اجازه استفاده از این فرم را ندارید')
        return redirect('site_profile:page_404')

    user = get_object_or_404(User, id=user_id)

    roles = Role.objects.filter(is_active=True)

    # نقش‌های فعلی کاربر
    user_roles = user.user_roles.filter(is_active=True).values_list('role__name', flat=True)
    user_roles_list = list(user_roles)

    if request.method == 'POST':
        first_name = request.POST.get('first_name')
        last_name = request.POST.get('last_name')
        mobile = request.POST.get('mobile')
        email = request.POST.get('email')
        is_active = request.POST.get('is_active') == 'on'
        selected_roles = request.POST.getlist('roles')
        new_password = request.POST.get('new_password')
        confirm_password = request.POST.get('confirm_password')

        errors = []

        if mobile and User.objects.filter(mobile=mobile).exclude(id=user.id).exists():
            errors.append('این شماره موبایل قبلاً ثبت شده است')

        if email and User.objects.filter(email=email).exclude(id=user.id).exists():
            errors.append('این ایمیل قبلاً ثبت شده است')

        if new_password:
            if len(new_password) < 4:
                errors.append('رمز عبور باید حداقل ۴ کاراکتر باشد')
            elif new_password != confirm_password:
                errors.append('رمز عبور و تکرار آن مطابقت ندارند')

        if errors:
            for error in errors:
                messages.error(request, error)
            return render(request, 'users/user_edit.html', {
                'user': user,
                'roles': roles,
                'user_roles': user_roles_list,
                'userId': request.user.id,
                'this_user': request.user,
            })

        try:
            user.first_name = first_name
            user.last_name = last_name
            user.mobile = mobile
            user.email = email
            user.is_active = is_active

            if new_password:
                user.password = make_password(new_password)

            user.save()

            # بروزرسانی نقش‌ها
            UserRole.objects.filter(user=user, is_active=True).update(is_active=False)

            for role_name in selected_roles:
                try:
                    role = Role.objects.get(name=role_name)
                    user_role, created = UserRole.objects.get_or_create(
                        user=user,
                        role=role,
                        defaults={'assigned_by': request.user, 'is_active': True}
                    )
                    if not created:
                        user_role.is_active = True
                        user_role.assigned_by = request.user
                        user_role.save()
                except Role.DoesNotExist:
                    pass

            messages.success(request, f'کاربر "{user.get_full_name()}" با موفقیت بروزرسانی شد')
            return redirect('site_profile:user_list')

        except Exception as e:
            messages.error(request, f'خطا در بروزرسانی کاربر: {str(e)}')

    # ✅ اضافه کردن تاریخ شمسی
    user.date_joined_shamsi = get_shamsi_datetime(user.date_joined)
    user.last_login_shamsi = get_shamsi_datetime(user.last_login)

    context = {
        'user': user,
        'roles': roles,
        'user_roles': user_roles_list,
        'title': 'ویرایش کاربر',
        'userId': request.user.id,
        'this_user': request.user,
        "perimissin_list": perimissin_list,
        "perimissin_group": perimissin_group,
    }
    return render(request, 'users/user_edit.html', context)


@login_required
def user_delete(request, user_id):
    """حذف کاربر (غیرفعال کردن)"""
    user = request.user
    permission, perimissin_list, perimissin_group = views_permissions(user, "user_delete")

    if permission is False:
        messages.error(request, 'شما اجازه استفاده از این فرم را ندارید')
        return redirect('site_profile:page_404')

    user = get_object_or_404(User, id=user_id)

    # جلوگیری از حذف خود کاربر
    if user.id == request.user.id:
        messages.error(request, 'شما نمی‌توانید خودتان را حذف کنید')
        return redirect('site_profile:user_list')

    # غیرفعال کردن به جای حذف
    user.is_active = False
    user.save()

    messages.success(request, f'کاربر "{user.get_full_name()}" با موفقیت غیرفعال شد')
    return redirect('site_profile:user_list')


@login_required
def user_detail(request, user_id):
    """مشاهده جزئیات کاربر"""
    user = request.user
    permission, perimissin_list, perimissin_group = views_permissions(user, "user_detail")

    if permission is False:
        messages.error(request, 'شما اجازه استفاده از این فرم را ندارید')
        return redirect('site_profile:page_404')

    user = get_object_or_404(User, id=user_id)

    # اطلاعات نقش‌ها
    user_roles = user.user_roles.filter(is_active=True).select_related('role')

    # ✅ اضافه کردن تاریخ شمسی
    user.date_joined_shamsi = get_shamsi_datetime(user.date_joined)
    user.last_login_shamsi = get_shamsi_datetime(user.last_login)
    if user.birth_date:
        user.birth_date_shamsi = get_shamsi_date(user.birth_date)
    else:
        user.birth_date_shamsi = ''

    # ✅ اضافه کردن تاریخ شمسی به هر نقش
    for ur in user_roles:
        ur.assigned_at_shamsi = get_shamsi_datetime(ur.assigned_at)

    context = {
        'user': user,
        'user_roles': user_roles,
        'title': 'جزئیات کاربر',
        'userId': request.user.id,
        'this_user': request.user,
        "perimissin_list": perimissin_list,
        "perimissin_group": perimissin_group,
    }
    return render(request, 'users/user_detail.html', context)


# ==================== مدیریت نقش‌ها ====================
@login_required
def role_list(request):
    """لیست نقش‌ها"""
    user = request.user
    permission, perimissin_list, perimissin_group = views_permissions(user, "role_list")

    if permission is False:
        messages.error(request, 'شما اجازه استفاده از این فرم را ندارید')
        return redirect('site_profile:page_404')

    roles = Role.objects.all().order_by('name')

    for role in roles:
        role.user_count = UserRole.objects.filter(role=role, is_active=True).count()
        role.permission_count = RolePermission.objects.filter(role=role).count()
        # ✅ اضافه کردن تاریخ شمسی
        role.created_at_shamsi = get_shamsi_date(role.created_at)
        role.updated_at_shamsi = get_shamsi_date(role.updated_at)

    context = {
        'roles': roles,
        'title': 'لیست نقش‌ها',
        'userId': request.user.id,
        'this_user': request.user,
        "perimissin_list": perimissin_list,
        "perimissin_group": perimissin_group,
    }
    return render(request, 'roles/role_list.html', context)


@login_required
def role_create(request):
    """ایجاد نقش جدید"""
    user = request.user
    permission, perimissin_list, perimissin_group = views_permissions(user, "role_create")

    if permission is False:
        messages.error(request, 'شما اجازه استفاده از این فرم را ندارید')
        return redirect('site_profile:page_404')

    if request.method == 'POST':
        name = request.POST.get('name')
        title = request.POST.get('title')
        description = request.POST.get('description')

        if not name or not title:
            messages.error(request, 'نام و عنوان نقش الزامی است')
            return render(request, 'roles/role_create.html', {
                'userId': request.user.id,
                'this_user': request.user,
            })

        if Role.objects.filter(name=name).exists():
            messages.error(request, f'نقش با نام "{name}" قبلاً تعریف شده است')
            return render(request, 'roles/role_create.html', {
                'userId': request.user.id,
                'this_user': request.user,
            })

        try:
            role = Role.objects.create(
                name=name,
                title=title,
                description=description,
                is_active=True
            )
            messages.success(request, f'نقش "{role.title}" با موفقیت ایجاد شد')
            return redirect('site_profile:role_list')
        except Exception as e:
            messages.error(request, f'خطا در ایجاد نقش: {str(e)}')

    context = {
        'title': 'افزودن نقش جدید',
        'userId': request.user.id,
        'this_user': request.user,
        "perimissin_list": perimissin_list,
        "perimissin_group": perimissin_group,
    }
    return render(request, 'roles/role_create.html', context)


@login_required
def role_edit(request, role_id):
    """ویرایش نقش"""
    user = request.user
    permission, perimissin_list, perimissin_group = views_permissions(user, "role_edit")

    if permission is False:
        messages.error(request, 'شما اجازه استفاده از این فرم را ندارید')
        return redirect('site_profile:page_404')

    role = get_object_or_404(Role, id=role_id)

    if request.method == 'POST':
        title = request.POST.get('title')
        description = request.POST.get('description')
        is_active = request.POST.get('is_active') == 'on'

        if not title:
            messages.error(request, 'عنوان نقش الزامی است')
            return render(request, 'roles/role_create.html', {
                'role': role,
                'userId': request.user.id,
                'this_user': request.user,
            })

        role.title = title
        role.description = description
        role.is_active = is_active
        role.save()

        messages.success(request, f'نقش "{role.title}" با موفقیت بروزرسانی شد')
        return redirect('site_profile:role_list')

    # ✅ اضافه کردن تاریخ شمسی
    role.created_at_shamsi = get_shamsi_datetime(role.created_at)
    role.updated_at_shamsi = get_shamsi_datetime(role.updated_at)

    context = {
        'role': role,
        'title': 'ویرایش نقش',
        'userId': request.user.id,
        'this_user': request.user,
        "perimissin_list": perimissin_list,
        "perimissin_group": perimissin_group,
    }
    return render(request, 'roles/role_create.html', context)


@login_required
def role_delete(request, role_id):
    """حذف نقش"""
    user = request.user
    permission, perimissin_list, perimissin_group = views_permissions(user, "role_delete")

    if permission is False:
        messages.error(request, 'شما اجازه استفاده از این فرم را ندارید')
        return redirect('site_profile:page_404')

    role = get_object_or_404(Role, id=role_id)

    # بررسی اینکه نقش به کاربر متصل نباشد
    if UserRole.objects.filter(role=role, is_active=True).exists():
        # به جای حذف، غیرفعال می‌کنیم
        role.is_active = False
        role.save()
        messages.warning(request, f'نقش "{role.title}" غیرفعال شد (چون به کاربر متصل است)')
    else:
        role.delete()
        messages.success(request, f'نقش "{role.title}" با موفقیت حذف شد')

    return redirect('site_profile:role_list')


@login_required
def role_permissions(request, role_id):
    """مدیریت مجوزهای یک نقش"""
    user = request.user
    permission, perimissin_list, perimissin_group = views_permissions(user, "role_permissions")

    if permission is False:
        messages.error(request, 'شما اجازه استفاده از این فرم را ندارید')
        return redirect('site_profile:page_404')

    role = get_object_or_404(Role, id=role_id)

    # مرتب‌سازی بر اساس group و سپس name
    all_permissions = Permission.objects.all().order_by('group', 'name')

    # مجوزهای فعلی نقش
    role_permissions = RolePermission.objects.filter(role=role).values_list('permission__id', flat=True)

    if request.method == 'POST':
        selected_permissions = request.POST.getlist('permissions')

        # حذف مجوزهای قبلی
        RolePermission.objects.filter(role=role).delete()

        # اضافه کردن مجوزهای جدید
        for perm_id in selected_permissions:
            try:
                permission = Permission.objects.get(id=perm_id)
                RolePermission.objects.create(role=role, permission=permission)
            except Permission.DoesNotExist:
                pass

        messages.success(request, f'مجوزهای نقش "{role.title}" با موفقیت بروزرسانی شد')
        return redirect('site_profile:role_list')

    context = {
        'role': role,
        'all_permissions': all_permissions,
        'role_permissions': list(role_permissions),
        'title': f'مدیریت مجوزهای {role.title}',
        'userId': request.user.id,
        'this_user': request.user,
        "perimissin_list": perimissin_list,
        "perimissin_group": perimissin_group,
    }
    return render(request, 'roles/role_permissions.html', context)


# ==================== مدیریت مجوزها ====================
@login_required
def permission_list(request):
    """لیست مجوزها"""
    user = request.user
    permission, perimissin_list, perimissin_group = views_permissions(user, "permission_list")

    if permission is False:
        messages.error(request, 'شما اجازه استفاده از این فرم را ندارید')
        return redirect('site_profile:page_404')

    permissions = Permission.objects.all().order_by('group', 'name')

    context = {
        'permissions': permissions,
        'title': 'لیست مجوزها',
        'userId': request.user.id,
        'this_user': request.user,
        "perimissin_list": perimissin_list,
        "perimissin_group": perimissin_group,
    }
    return render(request, 'permissions/permission_list.html', context)


@login_required
def permission_create(request):
    """ایجاد مجوز جدید"""
    user = request.user
    permission, perimissin_list, perimissin_group = views_permissions(user, "permission_create")

    if permission is False:
        messages.error(request, 'شما اجازه استفاده از این فرم را ندارید')
        return redirect('site_profile:page_404')

    if request.method == 'POST':
        name = request.POST.get('name')
        code = request.POST.get('code')
        description = request.POST.get('description')

        if not name or not code:
            messages.error(request, 'نام و کد مجوز الزامی است')
            return render(request, 'permissions/permission_create.html', {
                'userId': request.user.id,
                'this_user': request.user,
            })

        if Permission.objects.filter(code=code).exists():
            messages.error(request, f'مجوز با کد "{code}" قبلاً تعریف شده است')
            return render(request, 'permissions/permission_create.html', {
                'userId': request.user.id,
                'this_user': request.user,
            })

        try:
            permission = Permission.objects.create(
                name=name,
                code=code,
                description=description
            )
            messages.success(request, f'مجوز "{permission.name}" با موفقیت ایجاد شد')
            return redirect('site_profile:permission_list')
        except Exception as e:
            messages.error(request, f'خطا در ایجاد مجوز: {str(e)}')

    context = {
        'title': 'افزودن مجوز جدید',
        'userId': request.user.id,
        'this_user': request.user,
        "perimissin_list": perimissin_list,
        "perimissin_group": perimissin_group,
    }
    return render(request, 'permissions/permission_create.html', context)


@login_required
def permission_delete(request, permission_id):
    """حذف مجوز"""
    user = request.user
    permission, perimissin_list, perimissin_group = views_permissions(user, "permission_delete")

    if permission is False:
        messages.error(request, 'شما اجازه استفاده از این فرم را ندارید')
        return redirect('site_profile:page_404')

    permission = get_object_or_404(Permission, id=permission_id)

    # بررسی اینکه مجوز به نقشی متصل نباشد
    if RolePermission.objects.filter(permission=permission).exists():
        messages.warning(request, f'مجوز "{permission.name}" به نقش متصل است و قابل حذف نیست')
    else:
        permission.delete()
        messages.success(request, f'مجوز "{permission.name}" با موفقیت حذف شد')

    return redirect('site_profile:permission_list')


@login_required
def export_permissions_excel(request):
    """
    خروجی Excel از تمام مجوزها
    """
    user = request.user
    permission, perimissin_list, perimissin_group = views_permissions(user, "permission_list")
    if not permission:
        messages.error(request, 'شما دسترسی به این بخش را ندارید')
        return redirect('site_profile:page_404')

    permissions = Permission.objects.all().order_by('group', 'name')

    wb = Workbook()
    ws = wb.active
    ws.title = 'مجوزها'

    # هدرها
    headers = ['گروه', 'نام مجوز', 'کد', 'توضیحات']
    header_fill = PatternFill(start_color='366092', end_color='366092', fill_type='solid')
    header_font = Font(color='FFFFFF', bold=True, size=11)

    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal='center', vertical='center')

    # داده‌ها
    row_num = 2
    for perm in permissions:
        ws.cell(row=row_num, column=1, value=perm.group or '')
        ws.cell(row=row_num, column=2, value=perm.name)
        ws.cell(row=row_num, column=3, value=perm.code)
        ws.cell(row=row_num, column=4, value=perm.description or '')
        for col in range(1, 5):
            ws.cell(row=row_num, column=col).alignment = Alignment(horizontal='center', vertical='center')
        row_num += 1

    # تنظیم عرض ستون‌ها
    ws.column_dimensions['A'].width = 30
    ws.column_dimensions['B'].width = 35
    ws.column_dimensions['C'].width = 25
    ws.column_dimensions['D'].width = 50

    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    # ✅ نام فایل با تاریخ شمسی
    today_shamsi = get_shamsi_date(timezone.now()).replace('/', '-')
    response['Content-Disposition'] = f'attachment; filename=permissions_{today_shamsi}.xlsx'
    wb.save(response)
    return response


@login_required
def import_permissions_excel(request):
    """
    ورود مجوزها از فایل Excel
    """
    user = request.user
    permission, perimissin_list, perimissin_group = views_permissions(user, "permission_create")
    if not permission:
        messages.error(request, 'شما دسترسی به این بخش را ندارید')
        return redirect('site_profile:page_404')

    if request.method == 'POST' and request.FILES.get('excel_file'):
        excel_file = request.FILES['excel_file']
        try:
            wb = openpyxl.load_workbook(excel_file)
            ws = wb.active

            created_count = 0
            skipped_count = 0
            errors = []

            # شروع از سطر دوم (ردیف اول هدر است)
            for row in ws.iter_rows(min_row=2, values_only=True):
                group = row[0] if row[0] else ''
                name = row[1] if row[1] else ''
                code = row[2] if row[2] else ''
                description = row[3] if row[3] else ''

                # اعتبارسنجی ساده
                if not name or not code:
                    errors.append(f'ردیف {row[0]}: نام یا کد مجوز خالی است')
                    continue

                # بررسی تکراری بودن کد
                if Permission.objects.filter(code=code).exists():
                    skipped_count += 1
                    continue

                # ایجاد مجوز جدید
                Permission.objects.create(
                    group=group,
                    name=name,
                    code=code,
                    description=description
                )
                created_count += 1

            messages.success(
                request,
                f'✅ {created_count} مجوز جدید ایجاد شد. '
                f'{skipped_count} مجوز به دلیل تکراری بودن کد نادیده گرفته شدند.'
            )
            if errors:
                messages.warning(request, '⚠️ برخی ردیف‌ها دارای خطا بودند: ' + ' | '.join(errors[:5]))

        except Exception as e:
            messages.error(request, f'خطا در خواندن فایل: {str(e)}')
        return redirect('site_profile:permission_list')

    context = {
        'title': 'ورود مجوزها از Excel',
        'this_user': request.user,
        'perimissin_list': perimissin_list,
        'perimissin_group': perimissin_group,
    }
    return render(request, 'permissions/import_permissions.html', context)


@login_required
def download_permissions_template(request):
    """
    دانلود فایل نمونه Excel برای ورود مجوزها
    """
    wb = Workbook()
    ws = wb.active
    ws.title = 'مجوزها'

    headers = ['گروه', 'نام مجوز', 'کد', 'توضیحات']
    header_fill = PatternFill(start_color='366092', end_color='366092', fill_type='solid')
    header_font = Font(color='FFFFFF', bold=True, size=11)

    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal='center', vertical='center')

    # نمونه داده
    sample_data = [
        ['مدیریت کاربران', 'مشاهده کاربران', 'user.view', 'دسترسی به لیست کاربران'],
        ['مدیریت کاربران', 'ایجاد کاربر', 'user.create', 'امکان ایجاد کاربر جدید'],
        ['گزارشات', 'مشاهده گزارش مشتریان', 'report.customer', 'دسترسی به گزارش مشتریان'],
    ]
    for i, row in enumerate(sample_data, 2):
        for j, val in enumerate(row, 1):
            ws.cell(row=i, column=j, value=val)

    ws.column_dimensions['A'].width = 30
    ws.column_dimensions['B'].width = 35
    ws.column_dimensions['C'].width = 25
    ws.column_dimensions['D'].width = 50

    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = 'attachment; filename=template_import_permissions.xlsx'
    wb.save(response)
    return response