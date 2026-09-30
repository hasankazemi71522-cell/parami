# estate/utils_location.py

import openpyxl
import pandas as pd
from io import BytesIO
from django.core.files.uploadedfile import InMemoryUploadedFile


def read_excel_file(file):
    """
    خواندن فایل اکسل و برگرداندن دیتافریم پانداس
    """
    try:
        # خواندن با پانداس
        df = pd.read_excel(file)
        return df
    except Exception as e:
        raise ValueError(f"خطا در خواندن فایل اکسل: {str(e)}")


def validate_location_excel(df):
    """
    اعتبارسنجی ستون‌های فایل اکسل استان و شهرستان و محله
    """
    required_columns = ['استان', 'شهرستان', 'محله']

    # بررسی وجود ستون‌های مورد نیاز
    for col in required_columns:
        if col not in df.columns:
            raise ValueError(f"ستون '{col}' در فایل اکسل وجود ندارد")

    # حذف ردیف‌های خالی
    df = df.dropna(subset=['استان'])

    # استان‌ها نباید خالی باشند
    if df['استان'].isnull().any():
        raise ValueError("ستون 'استان' نمی‌تواند خالی باشد")

    return df


def process_location_data(df):
    """
    پردازش داده‌های اکسل و تبدیل به ساختار استان-شهرستان-محله
    """
    result = {}

    for index, row in df.iterrows():
        province_name = str(row['استان']).strip()
        city_name = str(row['شهرستان']).strip() if pd.notna(row['شهرستان']) else None
        neighborhood_name = str(row['محله']).strip() if pd.notna(row['محله']) else None

        # اگر شهرستان وجود نداشته باشد، از استان استفاده کن
        if not city_name:
            city_name = province_name

        # ساختار دیتا
        if province_name not in result:
            result[province_name] = {
                'cities': {}
            }

        if city_name not in result[province_name]['cities']:
            result[province_name]['cities'][city_name] = {
                'neighborhoods': []
            }

        if neighborhood_name and neighborhood_name.strip():
            result[province_name]['cities'][city_name]['neighborhoods'].append(neighborhood_name)

    return result


def generate_excel_template():
    """
    تولید فایل اکسل نمونه برای دانلود
    """
    import openpyxl
    from openpyxl.styles import Font, Alignment, PatternFill

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "مکان‌ها"

    # هدرها
    headers = ['استان', 'شهرستان', 'محله']
    header_fill = PatternFill(start_color="4CAF50", end_color="4CAF50", fill_type="solid")

    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")

    # تنظیم عرض ستون‌ها
    ws.column_dimensions['A'].width = 30
    ws.column_dimensions['B'].width = 30
    ws.column_dimensions['C'].width = 30

    # داده‌های نمونه
    sample_data = [
        ['تهران', 'تهران', 'ونک'],
        ['تهران', 'تهران', 'تجریش'],
        ['تهران', 'ری', ''],
        ['اصفهان', 'اصفهان', ''],
        ['اصفهان', 'کاشان', ''],
        ['فارس', 'شیراز', ''],
    ]

    for row_idx, row_data in enumerate(sample_data, 2):
        for col_idx, value in enumerate(row_data, 1):
            ws.cell(row=row_idx, column=col_idx, value=value)

    return wb